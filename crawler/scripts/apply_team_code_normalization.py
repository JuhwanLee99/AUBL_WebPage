#!/usr/bin/env python3
"""Normalize TEAM inserts/setters in generated import SQL.

This script:
1) Rewrites team-name based TEAM/SET statements to stable team_code statements.
2) Optionally ingests playoff workbook data (2015~2025) and appends normalized
   playoff tier results (으뜸/버금) into a dedicated SQL table.
"""

from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


PAT_NAME_INSERT = re.compile(
    r"^INSERT INTO TEAM \(team_name, team_code\) SELECT '((?:[^'\\]|\\.)*)', NULL "
    r"FROM DUAL WHERE NOT EXISTS \(SELECT 1 FROM TEAM WHERE team_name = '((?:[^'\\]|\\.)*)'\);$"
)
PAT_CODE_INSERT = re.compile(
    r"^INSERT INTO TEAM \(team_name, team_code\) SELECT '((?:[^'\\]|\\.)*)', '((?:[^'\\]|\\.)*)' "
    r"FROM DUAL WHERE NOT EXISTS \(SELECT 1 FROM TEAM WHERE team_code = '((?:[^'\\]|\\.)*)'\);$"
)
PAT_SET_NAME = re.compile(
    r"^SET (@[A-Za-z0-9_]+) = \(SELECT team_id FROM TEAM WHERE team_name = '((?:[^'\\]|\\.)*)' LIMIT 1\);$"
)
PAT_TEAM_NAME_COND = re.compile(r"(t\.team_name\s*=\s*)'((?:[^'\\]|\\.)*)'")


@dataclass(frozen=True)
class TeamRule:
    target_code: str
    target_name: str
    action: str
    note: str


@dataclass(frozen=True)
class PlayoffEntry:
    year: int
    team_name: str
    playoff_tier: str
    playoff_round: str
    source_text: str


TEAM_RULES: dict[str, TeamRule] = {
    # Merge to existing coded teams.
    "건국대서울 불소야구부": TeamRule("11866", "건국대(서울) 불소야구", "MERGE", "name variant"),
    "건국대(서울) 불소야구부": TeamRule("11866", "건국대(서울) 불소야구", "MERGE", "name variant"),
    "중앙대학교(서울) 랑데뷰": TeamRule("24016", "중앙대학교 랑데뷰", "MERGE", "name variant"),
    "단국대Pandas": TeamRule("21584", "단국대 PANDAS", "MERGE", "spacing/case variant"),
    "서울시립대학교Falcons": TeamRule("11709", "서울시립대학교FALCONS", "MERGE", "case variant"),
    "한양대ERICA H.I.B.A": TeamRule("24012", "한양대ERICA HIBA", "MERGE", "punctuation variant"),
    "한국산업기술대학교 WINNERS": TeamRule("6205", "한국공학대학교 WINNERS", "MERGE", "school rename variant"),
    "한성대학교 Power turtles": TeamRule("24010", "한성대학교 TURTLES", "MERGE", "name variant"),
    "한성대 TURTLES": TeamRule("24010", "한성대학교 TURTLES", "MERGE", "abbreviation variant"),
    "항공대 야구부": TeamRule("23593", "한국항공대 Astros", "MERGE", "alias variant"),
    "타키온즈": TeamRule("23969", "강남대학교 타키온즈", "MERGE", "alias variant"),
    # Keep historically distinct teams separated (do not merge to another club).
    "건국대 글로컬 Panthers": TeamRule("28867", "건국대 글로컬 Panthers", "NEW_CODE", "distinct team"),
    "서울과학기술대 HEROES": TeamRule("23987", "서울과학기술대 HEROES", "NEW_CODE", "distinct team"),
    "KNSU 루나틱스": TeamRule("11853", "KNSU 루나틱스", "NEW_CODE", "distinct team"),
    "KNSU한국체대 야구동아리": TeamRule("11853", "KNSU 루나틱스", "MERGE", "alias of KNSU 루나틱스"),
    # New codes where no reliable existing code exists.
    "TEAM MAZOR": TeamRule("90004", "TEAM MAZOR", "NEW_CODE", "canonical for MAZOR variants"),
    "2024 AUBL 올스타전(ROOKIE-MAZOR)": TeamRule("90004", "TEAM MAZOR", "MERGE", "merged into MAZOR canonical"),
    "2024AUBL 올스타전 TEAM MAZOR": TeamRule("90004", "TEAM MAZOR", "MERGE", "merged into MAZOR canonical"),
    "TEAM WILSON": TeamRule("90005", "TEAM WILSON", "NEW_CODE", "canonical for WILSON variants"),
    "2024 AUBL 올스타전 TEAM WILSON": TeamRule("90005", "TEAM WILSON", "MERGE", "merged into WILSON canonical"),
    "2024 AUBL 올스타전(ROOKIE-WILSON)": TeamRule("90005", "TEAM WILSON", "MERGE", "merged into WILSON canonical"),
    "2025 AUBL 올스타 Team ENOUGH": TeamRule("90006", "2025 AUBL 올스타 Team ENOUGH", "NEW_CODE", "distinct all-star side"),
    "2025 AUBL 올스타전 Team AXEL": TeamRule("90007", "2025 AUBL 올스타전 Team AXEL", "NEW_CODE", "distinct all-star side"),
}


# Workbook alias -> normalized TEAM.team_name in generated SQL.
PLAYOFF_TEAM_ALIASES: dict[str, str] = {
    "가천대": "가천 WIND",
    "가톨릭대": "가톨릭대학교 텀블러즈",
    "강남대": "강남대학교 타키온즈",
    "건국대글로컬": "건국대 팬서스",
    "건국대서울": "건국대(서울) 불소야구",
    "경기대": "경기대학교 KGB",
    "경희대국제": "경희대국제 LIONS",
    "경희대서울": "경희대학교(서울) BRAVES",
    "고려대": "고려대학교 백구회",
    "광운대": "광운대학교 페가수스",
    "국민대": "국민대학교 윈드밀스",
    "단국대죽전": "단국대 PANDAS",
    "단국대천안": "단국대학교 하운드",
    "동국대": "동국대학교 LAE",
    "명지대서울": "명지대학교(서울) 나이너스",
    "명지대용인": "명지대학교(용인) 아마야구 퍼펙트",
    "백석대": "백석대학교 칼로스",
    "상명대": "상명대BUCKS",
    "서강대": "서강대학교 야구반 알바트로스",
    "서경대": "서경대학교 적시타",
    "서일대": "서일대학교 맥나이츠(MAC Knights)",
    "서울과학기술대": "서울과학기술대 미르",
    "서울대": "서울대학교",
    "서울시립대": "서울시립대학교FALCONS",
    "성균관대": "성균관대학교 킹고야구반",
    "세종대": "세종대학교 세종킹스",
    "숭실대": "숭실대학교 oners",
    "아주대": "아주대학교 ABBA",
    "연세대": "연세대학교 EAGLES",
    "용인대": "용인대학교 백호야구단",
    "육군사관학교": "육군사관학교 Mulacs",
    "인천대": "인천대학교 바이킹",
    "인하대": "인하대학교 비룡",
    "중앙대서울": "중앙대학교 랑데뷰",
    "중앙대안성": "중앙대학교 지킴이",
    "한국공학대": "한국공학대학교 WINNERS",
    "한국교통대": "한국교통대학교 스윙스",
    "한국외국어대글로벌": "외대(글로벌) 유니온",
    "한국외국어대서울": "한국외대(서울) 야구부",
    "한국항공대": "한국항공대 Astros",
    "한성대": "한성대학교 TURTLES",
    "한신대": "한신대학교 갱스터",
    "한양대서울": "한양대학교 불새",
    "한양대에리카": "한양대ERICA HIBA",
    "한체대": "한국체대 루나틱스",
    "홍익대": "홍익대학교 위너스",
}


def _unesc(value: str) -> str:
    return value.replace("\\'", "'").replace("\\\\", "\\")


def _esc(value: str) -> str:
    return value.replace("\\", "\\\\").replace("'", "\\'")


def _write_mapping_csv(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["source_team_name", "target_team_name", "target_team_code", "action", "note"])
        for source in sorted(TEAM_RULES):
            rule = TEAM_RULES[source]
            writer.writerow([source, rule.target_name, rule.target_code, rule.action, rule.note])


def _normalize_for_match(value: str) -> str:
    text = value.strip().lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[()\-_/]", "", text)
    text = text.replace("대학교", "대")
    text = text.replace("대학", "대")
    text = text.replace("외국어대", "외대")
    text = text.replace("에리카", "erica")
    text = text.replace("서울)", "서울")
    text = re.sub(r"[^0-9a-z가-힣]", "", text)
    return text


def _extract_sql_team_names(sql_path: Path) -> set[str]:
    names: set[str] = set()
    with sql_path.open("r", encoding="utf-8") as handle:
        for raw in handle:
            line = raw.rstrip("\n")
            m_code = PAT_NAME_INSERT.match(line)
            if m_code:
                names.add(_unesc(m_code.group(1)))
                continue
            m_named = PAT_CODE_INSERT.match(line)
            if m_named:
                names.add(_unesc(m_named.group(1)))
    return names


def _resolve_playoff_team_name(source_team_name: str, sql_team_names: set[str]) -> str | None:
    candidate = PLAYOFF_TEAM_ALIASES.get(source_team_name, source_team_name)
    rule = TEAM_RULES.get(candidate)
    if rule is not None:
        candidate = rule.target_name
    if candidate in sql_team_names:
        return candidate

    norm_target = _normalize_for_match(candidate)
    exact = [name for name in sql_team_names if _normalize_for_match(name) == norm_target]
    if len(exact) == 1:
        return exact[0]

    contains = [
        name
        for name in sql_team_names
        if norm_target and (
            norm_target in _normalize_for_match(name) or _normalize_for_match(name) in norm_target
        )
    ]
    if len(contains) == 1:
        return contains[0]

    return None


def _extract_playoff_tier_and_round(source_text: str) -> tuple[str, str] | None:
    raw = source_text.strip()
    if not raw:
        return None
    compact = re.sub(r"\s+", "", raw).lstrip("*")
    if "으뜸" in compact and "버금" in compact:
        return None
    tier: str | None = None
    marker = ""
    if "으뜸" in compact:
        tier = "EUTTEUM"
        marker = "으뜸"
    elif "버금" in compact:
        tier = "BEOGEUM"
        marker = "버금"
    if tier is None:
        return None

    playoff_round_raw = compact.replace(marker, "").replace("*", "")
    standard_round = None
    for token in ("준우승", "우승", "16강", "8강", "4강"):
        if token in playoff_round_raw:
            standard_round = token
            break
    playoff_round = standard_round or (playoff_round_raw or "참가")
    return tier, playoff_round


def _iter_playoff_entries(
    playoff_xlsx: Path,
    sheet_name: str,
    start_year: int,
    end_year: int,
    sql_team_names: set[str],
) -> tuple[list[PlayoffEntry], list[tuple[int, str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - env dependent
        raise RuntimeError("openpyxl is required for --playoff-xlsx") from exc

    wb = load_workbook(playoff_xlsx, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    ws = wb[sheet_name]

    year_cols: dict[int, int] = {}
    for col in range(2, ws.max_column + 1):
        value = ws.cell(1, col).value
        year: int | None = None
        if isinstance(value, int):
            year = value
        elif isinstance(value, str):
            m = re.search(r"(20\d{2}|19\d{2})", value)
            if m:
                year = int(m.group(1))
        if year is None:
            continue
        if start_year <= year <= end_year:
            year_cols[col] = year

    entries: list[PlayoffEntry] = []
    unresolved: list[tuple[int, str, str]] = []
    dedupe: set[tuple[int, str, str]] = set()

    for row in range(2, ws.max_row + 1):
        raw_team = ws.cell(row, 1).value
        if not isinstance(raw_team, str):
            continue
        source_team = raw_team.strip()
        if not source_team:
            continue

        resolved_team = _resolve_playoff_team_name(source_team, sql_team_names)
        for col, year in year_cols.items():
            raw_result = ws.cell(row, col).value
            if not isinstance(raw_result, str):
                continue
            tier_info = _extract_playoff_tier_and_round(raw_result)
            if tier_info is None:
                continue
            playoff_tier, playoff_round = tier_info
            if resolved_team is None:
                unresolved.append((year, source_team, raw_result.strip()))
                continue
            key = (year, resolved_team, playoff_tier)
            if key in dedupe:
                continue
            dedupe.add(key)
            entries.append(
                PlayoffEntry(
                    year=year,
                    team_name=resolved_team,
                    playoff_tier=playoff_tier,
                    playoff_round=playoff_round,
                    source_text=raw_result.strip(),
                )
            )

    entries.sort(key=lambda e: (e.year, e.team_name, e.playoff_tier))
    unresolved.sort(key=lambda r: (r[0], r[1], r[2]))
    return entries, unresolved


def _append_playoff_sql(output_sql: Path, entries: Iterable[PlayoffEntry]) -> int:
    written = 0
    with output_sql.open("a", encoding="utf-8") as out:
        out.write("\n-- Playoff Tier Normalization (from workbook)\n")
        out.write(
            "CREATE TABLE IF NOT EXISTS PLAYOFF_TEAM_RESULT ("
            "ptr_id INT AUTO_INCREMENT PRIMARY KEY, "
            "season_id INT NOT NULL, "
            "team_id INT NOT NULL, "
            "playoff_tier VARCHAR(20) NOT NULL, "
            "playoff_round VARCHAR(30) NOT NULL, "
            "source_text VARCHAR(100) NULL, "
            "UNIQUE KEY uq_playoff_team_year_tier (season_id, team_id, playoff_tier)"
            ");\n"
        )
        for entry in entries:
            out.write(
                "INSERT INTO PLAYOFF_TEAM_RESULT "
                "(season_id, team_id, playoff_tier, playoff_round, source_text) "
                "SELECT s.season_id, t.team_id, "
                f"'{_esc(entry.playoff_tier)}', '{_esc(entry.playoff_round)}', '{_esc(entry.source_text)}' "
                "FROM SEASON s JOIN TEAM t ON t.team_name = "
                f"'{_esc(entry.team_name)}' "
                f"WHERE s.year = {entry.year} "
                "ON DUPLICATE KEY UPDATE "
                "playoff_round = VALUES(playoff_round), "
                "source_text = VALUES(source_text);\n"
            )
            written += 1
    return written


def _write_playoff_csv(path: Path, entries: Iterable[PlayoffEntry]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["year", "team_name", "playoff_tier", "playoff_round", "source_text"])
        for entry in entries:
            writer.writerow([entry.year, entry.team_name, entry.playoff_tier, entry.playoff_round, entry.source_text])


def _write_playoff_unresolved_csv(path: Path, unresolved: Iterable[tuple[int, str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["year", "source_team_name", "source_text"])
        for row in unresolved:
            writer.writerow(list(row))


def normalize_sql(input_path: Path, output_path: Path) -> dict[str, int]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    insert_replaced = 0
    set_replaced = 0
    team_name_cond_replaced = 0

    with input_path.open("r", encoding="utf-8") as src, output_path.open("w", encoding="utf-8") as out:
        for raw in src:
            line = raw.rstrip("\n")

            match = PAT_NAME_INSERT.match(line)
            if match:
                team_name = _unesc(match.group(1))
                rule = TEAM_RULES.get(team_name)
                if rule is not None:
                    out_line = (
                        "INSERT INTO TEAM (team_name, team_code) "
                        f"SELECT '{_esc(rule.target_name)}', '{_esc(rule.target_code)}' "
                        f"FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM TEAM WHERE team_code = '{_esc(rule.target_code)}');"
                    )
                    out.write(out_line + "\n")
                    insert_replaced += 1
                    continue

            match = PAT_SET_NAME.match(line)
            if match:
                var_name = match.group(1)
                team_name = _unesc(match.group(2))
                rule = TEAM_RULES.get(team_name)
                if rule is not None:
                    out_line = (
                        f"SET {var_name} = (SELECT team_id FROM TEAM WHERE team_code = '{_esc(rule.target_code)}');"
                    )
                    out.write(out_line + "\n")
                    set_replaced += 1
                    continue

            # Keep stats joins valid after alias merge by rewriting team-name predicates.
            replaced_in_line = 0

            def _replace_cond(match: re.Match[str]) -> str:
                nonlocal replaced_in_line
                prefix = match.group(1)
                team_name = _unesc(match.group(2))
                rule = TEAM_RULES.get(team_name)
                if rule is None:
                    return match.group(0)
                replaced_in_line += 1
                return f"{prefix}'{_esc(rule.target_name)}'"

            rewritten = PAT_TEAM_NAME_COND.sub(_replace_cond, line)
            if replaced_in_line:
                team_name_cond_replaced += replaced_in_line
                out.write(rewritten + "\n")
            else:
                out.write(raw)

    return {
        "insert_replaced": insert_replaced,
        "set_replaced": set_replaced,
        "team_name_cond_replaced": team_name_cond_replaced,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize team codes in import SQL.")
    parser.add_argument("input_sql", help="Input SQL file path")
    parser.add_argument("output_sql", help="Output SQL file path")
    parser.add_argument(
        "--mapping-csv",
        default="team_code_normalization_map.csv",
        help="Write applied mapping CSV to this path",
    )
    parser.add_argument(
        "--playoff-xlsx",
        default="",
        help="Optional playoff workbook path. If set, append playoff tier SQL normalization.",
    )
    parser.add_argument(
        "--playoff-sheet",
        default="2004~2025 시즌 결과",
        help="Sheet name inside playoff workbook (default: 2004~2025 시즌 결과)",
    )
    parser.add_argument(
        "--playoff-start-year",
        type=int,
        default=2015,
        help="Start year for playoff extraction (default: 2015)",
    )
    parser.add_argument(
        "--playoff-end-year",
        type=int,
        default=2025,
        help="End year for playoff extraction (default: 2025)",
    )
    parser.add_argument(
        "--playoff-csv",
        default="",
        help="Optional output CSV for normalized playoff entries",
    )
    parser.add_argument(
        "--playoff-unresolved-csv",
        default="",
        help="Optional output CSV for unresolved playoff team mappings",
    )
    parser.add_argument(
        "--strict-playoff-map",
        action="store_true",
        help="Fail with exit code 2 when unresolved playoff team mappings exist",
    )
    args = parser.parse_args()

    input_path = Path(args.input_sql).resolve()
    output_path = Path(args.output_sql).resolve()
    mapping_csv = Path(args.mapping_csv)
    if not mapping_csv.is_absolute():
        mapping_csv = (Path.cwd() / mapping_csv).resolve()

    stats = normalize_sql(input_path, output_path)
    _write_mapping_csv(mapping_csv)

    playoff_written = 0
    unresolved: list[tuple[int, str, str]] = []
    if args.playoff_xlsx:
        playoff_xlsx = Path(args.playoff_xlsx)
        if not playoff_xlsx.is_absolute():
            playoff_xlsx = (Path.cwd() / playoff_xlsx).resolve()
        sql_team_names = _extract_sql_team_names(output_path)
        entries, unresolved = _iter_playoff_entries(
            playoff_xlsx=playoff_xlsx,
            sheet_name=args.playoff_sheet,
            start_year=args.playoff_start_year,
            end_year=args.playoff_end_year,
            sql_team_names=sql_team_names,
        )
        playoff_written = _append_playoff_sql(output_path, entries)

        if args.playoff_csv:
            playoff_csv = Path(args.playoff_csv)
            if not playoff_csv.is_absolute():
                playoff_csv = (Path.cwd() / playoff_csv).resolve()
            _write_playoff_csv(playoff_csv, entries)
        if args.playoff_unresolved_csv:
            unresolved_csv = Path(args.playoff_unresolved_csv)
            if not unresolved_csv.is_absolute():
                unresolved_csv = (Path.cwd() / unresolved_csv).resolve()
            _write_playoff_unresolved_csv(unresolved_csv, unresolved)

    print(f"Input SQL:  {input_path}")
    print(f"Output SQL: {output_path}")
    print(f"Mapping:    {mapping_csv}")
    print(f"Replaced TEAM inserts: {stats['insert_replaced']}")
    print(f"Replaced TEAM SETs:    {stats['set_replaced']}")
    print(f"Replaced team_name predicates: {stats['team_name_cond_replaced']}")
    print(f"Rule count:            {len(TEAM_RULES)}")
    if args.playoff_xlsx:
        print(f"Playoff rows appended: {playoff_written}")
        print(f"Playoff unresolved:    {len(unresolved)}")
        if unresolved:
            for year, team_name, source_text in unresolved[:20]:
                print(f"  - unresolved: year={year}, team={team_name}, source={source_text}")
            if len(unresolved) > 20:
                print(f"  ... {len(unresolved) - 20} more unresolved rows")
        if args.strict_playoff_map and unresolved:
            raise SystemExit(2)


if __name__ == "__main__":
    main()
