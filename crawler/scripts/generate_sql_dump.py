#!/usr/bin/env python3
"""Generate SQL INSERT statements from crawler JSONL output files.

Usage:
    python scripts/generate_sql_dump.py /tmp/aubl_crawl_test --year 2024 > import_2024.sql

    # With playoff tier data from Excel workbook:
    python scripts/generate_sql_dump.py /tmp/aubl_crawl_test --year 2024 \\
        --playoff-xlsx "../2015~2025 플레이오프 정리.xlsx" > import_2024.sql

The generated SQL can be run on the server:
    mysql -u root webpage < import_2024.sql
"""
from __future__ import annotations

import json
import re
import sys
import zlib
from collections import Counter, defaultdict
from pathlib import Path

# ---------------------------------------------------------------------------
# Optional import from apply_team_code_normalization for playoff name matching
# ---------------------------------------------------------------------------
sys.path.insert(0, str(Path(__file__).parent))
try:
    from apply_team_code_normalization import (  # type: ignore[import]
        PLAYOFF_TEAM_ALIASES,
        _normalize_for_match as _playoff_normalize,
        _extract_playoff_tier_and_round,
    )
except Exception:
    def _playoff_normalize(value: str) -> str:  # type: ignore[misc]
        text = re.sub(r"\s+", "", value.strip().lower())
        text = re.sub(r"[()\-_/]", "", text)
        text = text.replace("대학교", "대").replace("대학", "대")
        return re.sub(r"[^0-9a-z가-힣]", "", text)

    PLAYOFF_TEAM_ALIASES: dict[str, str] = {}  # type: ignore[assignment]

    def _extract_playoff_tier_and_round(source_text: str):  # type: ignore[misc]
        return None


# Korean round text → DB enum value
_ROUND_MAP: dict[str, str] = {
    "16강": "ROUND_OF_16",
    "8강": "QUARTER_FINAL",
    "4강": "SEMI_FINAL",
    "준우승": "FINAL",
    "우승": "FINAL",
}


def _load_playoff_tier_map(xlsx_path: Path, sheet_name: str, year: int) -> dict[str, tuple[str, str]]:
    """Load team → (tier, round_korean) from playoff Excel for a given year.

    Returns a dict keyed by normalized team name.
    tier: 'EUTTEUM' | 'BEOGEUM'
    round_korean: '16강' | '8강' | '4강' | '준우승' | '우승'
    """
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError:
        print("WARNING: openpyxl not installed; playoff data skipped", file=sys.stderr)
        return {}

    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"WARNING: Sheet '{sheet_name}' not found; playoff data skipped", file=sys.stderr)
        return {}
    ws = wb[sheet_name]

    year_col = None
    for col in range(2, ws.max_column + 1):
        value = ws.cell(1, col).value
        if isinstance(value, int) and value == year:
            year_col = col
            break
        elif isinstance(value, str):
            m = re.search(r"(20\d{2}|19\d{2})", value)
            if m and int(m.group(1)) == year:
                year_col = col
                break

    if year_col is None:
        return {}

    result: dict[str, tuple[str, str]] = {}
    for row in range(2, ws.max_row + 1):
        raw_team = ws.cell(row, 1).value
        if not isinstance(raw_team, str):
            continue
        source_team = raw_team.strip()
        if not source_team:
            continue
        raw_result = ws.cell(row, year_col).value
        if not isinstance(raw_result, str):
            continue
        tier_info = _extract_playoff_tier_and_round(raw_result)
        if tier_info is None:
            continue
        tier, round_text = tier_info
        resolved = PLAYOFF_TEAM_ALIASES.get(source_team, source_team)
        norm = _playoff_normalize(resolved)
        if norm:
            result[norm] = (tier, round_text)

    return result


def _get_game_playoff_tier_and_round(
    home_name: str | None,
    away_name: str | None,
    home_score,
    away_score,
    tier_map: dict[str, tuple[str, str]],
) -> tuple[str | None, str | None]:
    """Determine playoff_tier and playoff_round for a single game.

    Uses the loser's best round from tier_map to determine the game round.
    Returns (playoff_tier, playoff_round) — either may be None.
    """
    if not tier_map:
        return None, None

    home_norm = _playoff_normalize(home_name or "")
    away_norm = _playoff_normalize(away_name or "")
    home_info = tier_map.get(home_norm)
    away_info = tier_map.get(away_norm)

    # Determine tier
    playoff_tier: str | None = None
    if home_info and away_info:
        playoff_tier = home_info[0] if home_info[0] == away_info[0] else home_info[0]
    elif home_info:
        playoff_tier = home_info[0]
    elif away_info:
        playoff_tier = away_info[0]

    # Determine round from loser's best Korean round text
    playoff_round: str | None = None
    try:
        hs = int(home_score) if home_score is not None else None
        as_ = int(away_score) if away_score is not None else None
    except (TypeError, ValueError):
        hs = as_ = None

    if hs is not None and as_ is not None and hs != as_:
        loser_info = away_info if hs > as_ else home_info
        if loser_info:
            playoff_round = _ROUND_MAP.get(loser_info[1])

    return playoff_tier, playoff_round


def _strip_jersey(name: str) -> str:
    """'김민혁(52)' → '김민혁', '김동혁 (91)' → '김동혁'"""
    return re.sub(r"\s*[\(\(]\d+[\)\)]\s*$", "", name).strip()


def _extract_jersey(name: str | None) -> int | None:
    """Extract trailing jersey number from player name."""
    if not isinstance(name, str):
        return None
    match = re.search(r"\(\s*(\d{1,3})\s*\)\s*$", name.strip())
    if not match:
        return None
    try:
        return int(match.group(1))
    except (TypeError, ValueError):
        return None


def _esc(val: str | None) -> str:
    if val is None:
        return "NULL"
    return "'" + val.replace("\\", "\\\\").replace("'", "\\'") + "'"


def _num(val) -> str:
    if val is None:
        return "NULL"
    return str(val)


def _parse_innings(val) -> str:
    """Parse '22 ⅔' → '22.7'"""
    if val is None:
        return "NULL"
    s = str(val).strip()
    frac_map = {"⅓": ".3", "⅔": ".7", "1/3": ".3", "2/3": ".7"}
    for frac, dec in frac_map.items():
        if frac in s:
            whole = s.replace(frac, "").strip()
            return (whole or "0") + dec
    try:
        return str(float(s))
    except (ValueError, TypeError):
        return "NULL"


def _normalize_team_name(name: str | None) -> str:
    if not isinstance(name, str):
        return ""
    compact = re.sub(r"\s+", "", name).lower()
    return re.sub(r"[^0-9a-z가-힣]", "", compact)


def _ctx_code(value) -> str | None:
    """Normalize crawler context code values: '0'/'-1'/blank -> None."""
    if value is None:
        return None
    text = str(value).strip()
    if text in {"", "0", "-1", "None", "none", "NULL", "null"}:
        return None
    return text


def _qualified_section(value) -> str | None:
    """Normalize section labels to IN/OUT when possible."""
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    compact = re.sub(r"\s+", "", raw).upper()
    if "OUT" in compact:
        return "OUT"
    if "IN" in compact:
        return "IN"
    return raw


def _synthetic_user_id(team_token: str, player_name: str) -> int:
    """Build deterministic negative user_id from team token + player name."""
    seed = f"{team_token}::{player_name}".encode("utf-8")
    value = zlib.crc32(seed) & 0x7FFFFFFF
    if value == 0:
        value = 1
    return -value


def _iter_jsonl(path: Path):
    if not path.exists():
        return
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            if isinstance(row, dict):
                yield row


def _row_year(row: dict) -> int | None:
    value = row.get("year")
    if isinstance(value, int):
        return value
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _latest_row_for_year(path: Path, year: int) -> dict | None:
    selected: dict | None = None
    for row in _iter_jsonl(path) or []:
        if _row_year(row) != year:
            continue
        selected = row
    return selected


def main():
    if len(sys.argv) < 2:
        print("Usage: python generate_sql_dump.py <jsonl_dir> [--year YEAR] [--playoff-xlsx PATH] [--playoff-sheet SHEET]", file=sys.stderr)
        sys.exit(1)

    data_dir = Path(sys.argv[1])
    year = 2024
    if "--year" in sys.argv:
        idx = sys.argv.index("--year")
        year = int(sys.argv[idx + 1])

    playoff_xlsx: Path | None = None
    if "--playoff-xlsx" in sys.argv:
        idx = sys.argv.index("--playoff-xlsx")
        playoff_xlsx = Path(sys.argv[idx + 1])
        if not playoff_xlsx.is_absolute():
            playoff_xlsx = (Path.cwd() / playoff_xlsx).resolve()

    playoff_sheet = "2004~2025 시즌 결과"
    if "--playoff-sheet" in sys.argv:
        idx = sys.argv.index("--playoff-sheet")
        playoff_sheet = sys.argv[idx + 1]

    # Load playoff tier map if Excel provided
    playoff_tier_map: dict[str, tuple[str, str]] = {}
    if playoff_xlsx and playoff_xlsx.exists():
        playoff_tier_map = _load_playoff_tier_map(playoff_xlsx, playoff_sheet, year)
        print(f"-- Loaded playoff tier map: {len(playoff_tier_map)} teams for year {year}", file=sys.stderr)
    elif playoff_xlsx:
        print(f"WARNING: --playoff-xlsx path not found: {playoff_xlsx}", file=sys.stderr)

    out = sys.stdout
    out.write("-- AUBL Crawler Data Import\n")
    out.write(f"-- Year: {year}\n")
    out.write("-- Generated by generate_sql_dump.py\n\n")
    out.write("SET NAMES utf8mb4;\n")
    out.write("SET FOREIGN_KEY_CHECKS = 0;\n\n")
    out.write("-- Ensure context columns exist (idempotent)\n")
    out.write("ALTER TABLE GAME ADD COLUMN IF NOT EXISTS league_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE GAME ADD COLUMN IF NOT EXISTS part_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_STATS ADD COLUMN IF NOT EXISTS league_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_STATS ADD COLUMN IF NOT EXISTS part_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_STATS ADD COLUMN IF NOT EXISTS qualified_section VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_STATS ADD COLUMN IF NOT EXISTS league_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_STATS ADD COLUMN IF NOT EXISTS part_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_STATS ADD COLUMN IF NOT EXISTS qualified_section VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_GAME_LOG ADD COLUMN IF NOT EXISTS game_type VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_GAME_LOG ADD COLUMN IF NOT EXISTS league_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE BATTER_GAME_LOG ADD COLUMN IF NOT EXISTS part_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_GAME_LOG ADD COLUMN IF NOT EXISTS game_type VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_GAME_LOG ADD COLUMN IF NOT EXISTS league_code VARCHAR(20) NULL;\n")
    out.write("ALTER TABLE PITCHER_GAME_LOG ADD COLUMN IF NOT EXISTS part_code VARCHAR(20) NULL;\n\n")
    out.write("ALTER TABLE TEAM_PLAYER ADD COLUMN IF NOT EXISTS jersey_number INT NULL;\n\n")

    # 1. SEASON
    out.write(f"-- Season\n")
    out.write(f"INSERT IGNORE INTO SEASON (year) VALUES ({year});\n")
    out.write(f"SET @season_id = (SELECT season_id FROM SEASON WHERE year = {year});\n\n")

    # 2. TEAM — collect unique teams
    teams_file = data_dir / "teams.jsonl"
    teams: dict[int, str] = {}  # team_idx → name
    team_id_map: dict[int, str] = {}  # team_idx → SQL variable name
    team_name_map: dict[str, str] = {}  # team_name → SQL variable name
    team_name_norm_map: dict[str, str] = {}  # normalized team_name → SQL variable name
    team_idx_by_name: dict[str, int] = {}
    team_idx_by_norm_name: dict[str, int] = {}
    for t in _iter_jsonl(teams_file) or []:
        if _row_year(t) != year:
            continue
        idx = t.get("team_idx")
        name = t.get("name")
        if idx and name:
            teams[idx] = name
            team_idx_by_name[name] = idx
            normalized_name = _normalize_team_name(name)
            if normalized_name:
                team_idx_by_norm_name[normalized_name] = idx

    out.write("-- Teams\n")
    for i, (idx, name) in enumerate(sorted(teams.items(), key=lambda x: x[0])):
        var = f"@team_{i}"
        team_id_map[idx] = var
        team_name_map[name] = var
        normalized_name = _normalize_team_name(name)
        if normalized_name:
            team_name_norm_map[normalized_name] = var
        out.write(
            f"INSERT INTO TEAM (team_name, team_code) "
            f"SELECT {_esc(name)}, {_esc(str(idx))} FROM DUAL "
            f"WHERE NOT EXISTS (SELECT 1 FROM TEAM WHERE team_code = {_esc(str(idx))});\n"
        )
        out.write(f"SET {var} = (SELECT team_id FROM TEAM WHERE team_code = {_esc(str(idx))});\n")
    out.write("\n")

    def _resolve_team_var(team_idx, team_name):
        var = team_id_map.get(team_idx)
        if var:
            return var
        if isinstance(team_name, str):
            var = team_name_map.get(team_name)
            if var:
                return var
            normalized = _normalize_team_name(team_name)
            if normalized:
                return team_name_norm_map.get(normalized)
        return None

    def _resolve_team_idx(team_idx, team_name) -> int | None:
        if isinstance(team_idx, int):
            return team_idx
        if isinstance(team_idx, str) and team_idx.strip().isdigit():
            return int(team_idx.strip())
        if isinstance(team_name, str):
            direct = team_idx_by_name.get(team_name)
            if direct is not None:
                return direct
            normalized = _normalize_team_name(team_name)
            if normalized:
                return team_idx_by_norm_name.get(normalized)
        return None

    def _player_user_id(player_name: str, team_idx=None, team_name=None) -> int:
        resolved_team_idx = _resolve_team_idx(team_idx, team_name)
        if resolved_team_idx is not None:
            team_token = f"T{resolved_team_idx}"
        elif isinstance(team_name, str) and team_name.strip():
            team_token = f"N{_normalize_team_name(team_name)}"
        else:
            team_token = "GLOBAL"
        return _synthetic_user_id(team_token, player_name)

    # 3. PLAYER — collect unique players, strip jersey numbers
    players_file = data_dir / "players.jsonl"
    players: dict[int, dict] = {}  # synthetic_user_id → {name, position}
    for p in _iter_jsonl(players_file) or []:
        name = _strip_jersey(p.get("name") or "")
        if not name:
            continue
        user_id = _player_user_id(name, team_idx=p.get("team_idx"), team_name=None)
        if user_id not in players:
            players[user_id] = {"name": name, "position": p.get("position")}

    out.write("-- Players\n")
    for user_id, info in sorted(players.items(), key=lambda x: x[0]):
        name = info.get("name")
        pos = info.get("position")
        out.write(
            f"INSERT INTO PLAYER (user_id, player_name, position, is_player) "
            f"SELECT {_num(user_id)}, {_esc(name)}, {_esc(pos)}, 1 FROM DUAL "
            f"WHERE NOT EXISTS (SELECT 1 FROM PLAYER WHERE user_id = {_num(user_id)});\n"
        )
    out.write("\n")

    # 4. TEAM_PLAYER — roster
    roster_file = data_dir / "roster_players.jsonl"
    out.write("-- Team-Player Roster\n")
    for r in _iter_jsonl(roster_file) or []:
        if _row_year(r) != year:
            continue
        team_idx = r.get("team_idx")
        raw_name = r.get("name") or ""
        name = _strip_jersey(raw_name)
        jersey = _extract_jersey(raw_name)
        user_id = _player_user_id(name, team_idx=team_idx, team_name=None)
        if not name or not team_idx:
            continue
        team_var = team_id_map.get(team_idx)
        if not team_var:
            continue
        player_lookup = f"(SELECT player_id FROM PLAYER WHERE user_id = {_num(user_id)} LIMIT 1)"
        out.write(
            f"INSERT INTO TEAM_PLAYER (team_id, player_id, season_id, jersey_number) "
            f"SELECT {team_var}, "
            f"{player_lookup}, "
            f"@season_id, {_num(jersey)} FROM DUAL "
            f"WHERE {player_lookup} IS NOT NULL "
            f"AND NOT EXISTS ("
            f"SELECT 1 FROM TEAM_PLAYER WHERE team_id = {team_var} "
            f"AND player_id = {player_lookup} "
            f"AND season_id = @season_id);\n"
        )
        if jersey is not None:
            out.write(
                f"UPDATE TEAM_PLAYER SET jersey_number = COALESCE(jersey_number, {_num(jersey)}) "
                f"WHERE team_id = {team_var} "
                f"AND player_id = {player_lookup} "
                f"AND season_id = @season_id;\n"
            )
    out.write("\n")

    # Build per-team context hints from league ranking rows so GAME can keep
    # league/part metadata even when schedule payload omits them.
    team_context_counter: dict[str, dict[str, Counter[str]]] = defaultdict(
        lambda: {"league": Counter(), "part": Counter()}
    )

    def _accumulate_team_context(team_name: str | None, league_code: str | None, part_code: str | None) -> None:
        if not isinstance(team_name, str) or not team_name.strip():
            return
        key = _normalize_team_name(team_name)
        if not key:
            return
        if league_code:
            team_context_counter[key]["league"][league_code] += 1
        if part_code:
            team_context_counter[key]["part"][part_code] += 1

    def _resolve_team_context(team_name: str | None) -> tuple[str | None, str | None]:
        if not isinstance(team_name, str) or not team_name.strip():
            return None, None
        key = _normalize_team_name(team_name)
        bucket = team_context_counter.get(key)
        if not bucket:
            return None, None
        league_counter = bucket["league"]
        part_counter = bucket["part"]
        league_code = league_counter.most_common(1)[0][0] if league_counter else None
        part_code = part_counter.most_common(1)[0][0] if part_counter else None
        return league_code, part_code

    # 5. BATTER_STATS from league records
    league_bat_file = data_dir / "league_batting_records.jsonl"
    out.write("-- Batter Stats (league records)\n")
    data = _latest_row_for_year(league_bat_file, year)
    if data:
        payload = data.get("payload", {})
        records = []
        d = payload.get("data")
        if isinstance(d, dict):
            records = d.get("records", [])
        elif isinstance(d, list):
            records = d

        for rec in records:
            raw_name = rec.get("mb_name") or rec.get("name") or ""
            name = _strip_jersey(raw_name)
            jersey = _extract_jersey(raw_name)
            team_name = rec.get("club_name") or rec.get("team_name") or ""
            if not name:
                continue
            player_user_id = _player_user_id(name, team_idx=None, team_name=team_name)
            team_var = _resolve_team_var(None, team_name)
            team_condition = f"tp.team_id = {team_var}" if team_var else f"t.team_name = {_esc(team_name)}"
            league_code = _ctx_code(rec.get("league_code") or rec.get("group_code"))
            part_code = _ctx_code(rec.get("part_code"))
            qualified_section = _qualified_section(rec.get("section"))
            _accumulate_team_context(team_name, league_code, part_code)

            gp = _num(rec.get("mygamecnt"))
            pa = _num(rec.get("bats"))
            ab = _num(rec.get("bat_cnt"))
            h = _num(rec.get("hit_cnt"))
            doubles = _num(rec.get("twobase"))
            triples = _num(rec.get("threebase"))
            hr = _num(rec.get("homerun"))
            rbi = _num(rec.get("bat_point"))
            sb = _num(rec.get("steal"))
            bb = _num(rec.get("fourball"))
            so = _num(rec.get("strikeout"))
            avg = _num(rec.get("hit_rate"))
            obp = _num(rec.get("base_rate"))
            slg = _num(rec.get("hitbase_rate"))
            ops = _num(rec.get("ops"))

            # League records are always regular season (season_type = NULL)
            out.write(
                f"INSERT INTO BATTER_STATS "
                f"(tp_id, season_id, games_played, plate_appearance, at_bats, hits, "
                f"doubles, triples, home_runs, runs_batted_in, stolen_bases, walks, strikeouts, "
                f"batting_average, on_base_pct, slugging_pct, ops, league_code, part_code, qualified_section) "
                f"SELECT tp.tp_id, @season_id, {gp}, {pa}, {ab}, {h}, "
                f"{doubles}, {triples}, {hr}, {rbi}, {sb}, {bb}, {so}, "
                f"{avg}, {obp}, {slg}, {ops}, {_esc(league_code)}, {_esc(part_code)}, {_esc(qualified_section)} "
                f"FROM TEAM_PLAYER tp "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"JOIN TEAM t ON tp.team_id = t.team_id "
                f"WHERE p.user_id = {_num(player_user_id)} "
                f"AND {team_condition} "
                f"AND tp.season_id = @season_id "
                f"AND NOT EXISTS ("
                f"SELECT 1 FROM BATTER_STATS bs WHERE bs.tp_id = tp.tp_id "
                f"AND bs.season_id = @season_id AND bs.season_type IS NULL) "
                f"LIMIT 1;\n"
            )
            if jersey is not None:
                out.write(
                    f"UPDATE TEAM_PLAYER tp "
                    f"JOIN PLAYER p ON tp.player_id = p.player_id "
                    f"JOIN TEAM t ON tp.team_id = t.team_id "
                    f"SET tp.jersey_number = COALESCE(tp.jersey_number, {_num(jersey)}) "
                    f"WHERE p.user_id = {_num(player_user_id)} "
                    f"AND {team_condition} "
                    f"AND tp.season_id = @season_id;\n"
                )
    out.write("\n")

    # 6. PITCHER_STATS from league records
    league_pitch_file = data_dir / "league_pitching_records.jsonl"
    out.write("-- Pitcher Stats (league records)\n")
    data = _latest_row_for_year(league_pitch_file, year)
    if data:
        payload = data.get("payload", {})
        records = []
        d = payload.get("data")
        if isinstance(d, dict):
            records = d.get("records", [])
        elif isinstance(d, list):
            records = d

        for rec in records:
            raw_name = rec.get("mb_name") or rec.get("name") or ""
            name = _strip_jersey(raw_name)
            jersey = _extract_jersey(raw_name)
            team_name = rec.get("club_name") or rec.get("team_name") or ""
            if not name:
                continue
            player_user_id = _player_user_id(name, team_idx=None, team_name=team_name)
            team_var = _resolve_team_var(None, team_name)
            team_condition = f"tp.team_id = {team_var}" if team_var else f"t.team_name = {_esc(team_name)}"
            league_code = _ctx_code(rec.get("league_code") or rec.get("group_code"))
            part_code = _ctx_code(rec.get("part_code"))
            qualified_section = _qualified_section(rec.get("section"))
            _accumulate_team_context(team_name, league_code, part_code)

            gp = _num(rec.get("mygamecnt"))
            ip = _parse_innings(rec.get("inning"))
            w = _num(rec.get("win"))
            l = _num(rec.get("lose"))
            sv = _num(rec.get("save"))
            hld = _num(rec.get("hold"))
            ha = _num(rec.get("nohit"))
            runs = _num(rec.get("lost_point"))
            er = _num(rec.get("self_point"))
            hra = _num(rec.get("nohomerun"))
            bb = _num(rec.get("fourball"))
            so = _num(rec.get("strikeout"))
            hbp = _num(rec.get("deadball"))
            wp = _num(rec.get("wildpitch"))
            bk = _num(rec.get("balk"))
            era = _num(rec.get("def_rate"))
            whip = _num(rec.get("whip"))
            k9 = _num(rec.get("strikeout_rate"))

            # League records are always regular season (season_type = NULL)
            out.write(
                f"INSERT INTO PITCHER_STATS "
                f"(tp_id, season_id, games_played, innings_pitched, wins, losses, saves, holds, "
                f"hits_allowed, runs_allowed, earned_runs, home_runs_allow, walks_allowed, strikeouts, "
                f"hit_batters, wild_pitches, balks, era, whip, k_per_9, league_code, part_code, qualified_section) "
                f"SELECT tp.tp_id, @season_id, {gp}, {ip}, {w}, {l}, {sv}, {hld}, "
                f"{ha}, {runs}, {er}, {hra}, {bb}, {so}, "
                f"{hbp}, {wp}, {bk}, {era}, {whip}, {k9}, {_esc(league_code)}, {_esc(part_code)}, {_esc(qualified_section)} "
                f"FROM TEAM_PLAYER tp "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"JOIN TEAM t ON tp.team_id = t.team_id "
                f"WHERE p.user_id = {_num(player_user_id)} "
                f"AND {team_condition} "
                f"AND tp.season_id = @season_id "
                f"AND NOT EXISTS ("
                f"SELECT 1 FROM PITCHER_STATS ps WHERE ps.tp_id = tp.tp_id "
                f"AND ps.season_id = @season_id AND ps.season_type IS NULL) "
                f"LIMIT 1;\n"
            )
            if jersey is not None:
                out.write(
                    f"UPDATE TEAM_PLAYER tp "
                    f"JOIN PLAYER p ON tp.player_id = p.player_id "
                    f"JOIN TEAM t ON tp.team_id = t.team_id "
                    f"SET tp.jersey_number = COALESCE(tp.jersey_number, {_num(jersey)}) "
                    f"WHERE p.user_id = {_num(player_user_id)} "
                    f"AND {team_condition} "
                    f"AND tp.season_id = @season_id;\n"
                )
    out.write("\n")

    # 7. GAME + BATTER_GAME_LOG + PITCHER_GAME_LOG from matches/batting_stats/pitching_stats
    matches_file = data_dir / "matches.jsonl"
    batting_file = data_dir / "batting_stats.jsonl"
    pitching_file = data_dir / "pitching_stats.jsonl"

    # Load matches
    matches: dict[int, dict] = {}
    for m in _iter_jsonl(matches_file) or []:
        if _row_year(m) != year:
            continue
        gidx = m.get("game_idx")
        if gidx:
            matches[gidx] = m

    # Some team names appear only in match headers (not in roster team list).
    # Register them by name so GAME home/away mapping does not collapse to one side.
    extra_team_names: list[str] = []
    seen_extra: set[str] = set()
    for match in matches.values():
        for key in ("home_team_name", "away_team_name"):
            name = match.get(key)
            if not isinstance(name, str) or not name.strip():
                continue
            normalized_name = _normalize_team_name(name)
            if name in team_name_map:
                continue
            if normalized_name and normalized_name in team_name_norm_map:
                continue
            if name in seen_extra:
                continue
            seen_extra.add(name)
            extra_team_names.append(name)

    if extra_team_names:
        out.write("-- Additional Teams inferred from matches\n")
        for i, name in enumerate(sorted(extra_team_names)):
            var = f"@team_extra_{i}"
            team_name_map[name] = var
            normalized_name = _normalize_team_name(name)
            if normalized_name:
                team_name_norm_map[normalized_name] = var
            out.write(
                f"INSERT INTO TEAM (team_name, team_code) "
                f"SELECT {_esc(name)}, NULL FROM DUAL "
                f"WHERE NOT EXISTS (SELECT 1 FROM TEAM WHERE team_name = {_esc(name)});\n"
            )
            out.write(f"SET {var} = (SELECT team_id FROM TEAM WHERE team_name = {_esc(name)} LIMIT 1);\n")
        out.write("\n")

    # Load batting stats by game_idx
    batting_by_game: dict[int, list] = defaultdict(list)
    for b in _iter_jsonl(batting_file) or []:
        if _row_year(b) != year:
            continue
        gidx = b.get("game_idx")
        if gidx:
            batting_by_game[gidx].append(b)

    # Load pitching stats by game_idx
    pitching_by_game: dict[int, list] = defaultdict(list)
    for p in _iter_jsonl(pitching_file) or []:
        if _row_year(p) != year:
            continue
        gidx = p.get("game_idx")
        if gidx:
            pitching_by_game[gidx].append(p)

    out.write("-- Games and Game Logs\n")
    for gidx, match in sorted(matches.items()):
        home_idx = match.get("home_team_idx")
        away_idx = match.get("away_team_idx")
        home_name = match.get("home_team_name")
        away_name = match.get("away_team_name")
        home_var = _resolve_team_var(home_idx, home_name)
        away_var = _resolve_team_var(away_idx, away_name)
        # Some source payloads have incorrect away team_idx. Fall back to name map.
        if home_var and away_var and home_var == away_var and home_name != away_name:
            resolved_home = _resolve_team_var(None, home_name)
            resolved_away = _resolve_team_var(None, away_name)
            if resolved_home:
                home_var = resolved_home
            if resolved_away:
                away_var = resolved_away
        home_var = home_var or "NULL"
        away_var = away_var or "NULL"
        home_score = _num(match.get("home_runs"))
        away_score = _num(match.get("away_runs"))
        league_code = _ctx_code(match.get("group_code"))
        part_code = _ctx_code(match.get("part_code"))
        if league_code is None or part_code is None:
            home_league_code, home_part_code = _resolve_team_context(home_name)
            away_league_code, away_part_code = _resolve_team_context(away_name)
            if league_code is None:
                league_code = home_league_code or away_league_code
            if part_code is None:
                part_code = home_part_code or away_part_code
        phase = match.get("phase")
        game_type = None
        playoff_tier: str | None = None
        playoff_round: str | None = None
        if isinstance(phase, str):
            normalized = phase.strip().lower()
            if normalized == "result":
                game_type = "정규시즌"
            elif normalized == "playoff":
                game_type = "포스트시즌"
                # Determine tier and round from Excel data
                playoff_tier, playoff_round = _get_game_playoff_tier_and_round(
                    home_name, away_name,
                    match.get("home_runs"), match.get("away_runs"),
                    playoff_tier_map,
                )
            elif normalized:
                game_type = normalized.upper()

        out.write(f"\n-- Game {gidx}\n")
        out.write(
            f"INSERT INTO GAME (season_id, game_date, game_number, home_team, away_team, "
            f"home_score, away_score, game_type, playoff_tier, playoff_round, league_code, part_code) "
            f"SELECT @season_id, CURDATE(), {gidx}, {home_var}, {away_var}, "
            f"{home_score}, {away_score}, {_esc(game_type)}, {_esc(playoff_tier)}, {_esc(playoff_round)}, "
            f"{_esc(league_code)}, {_esc(part_code)} "
            f"FROM DUAL WHERE NOT EXISTS ("
            f"SELECT 1 FROM GAME WHERE season_id = @season_id AND game_number = {gidx});\n"
        )
        out.write(f"SET @game_id = (SELECT game_id FROM GAME WHERE season_id = @season_id AND game_number = {gidx});\n")

        # Determine season_type for stats rows in this game
        # Regular season → NULL (no season_type), Playoff → tier code or 'PLAYOFF' fallback
        if game_type == "포스트시즌":
            stat_season_type: str | None = playoff_tier or "PLAYOFF"
        else:
            stat_season_type = None  # regular season stats: season_type IS NULL

        # Helper: season_type SQL literals for NOT EXISTS checks and INSERTs
        if stat_season_type is None:
            _st_esc = "NULL"
            _st_not_exists = "bs.season_type IS NULL"
            _ps_not_exists = "ps.season_type IS NULL"
        else:
            _st_esc = _esc(stat_season_type)
            _st_not_exists = f"bs.season_type = {_esc(stat_season_type)}"
            _ps_not_exists = f"ps.season_type = {_esc(stat_season_type)}"

        # Batter game logs — ensure TEAM_PLAYER + BATTER_STATS exist first
        for b in batting_by_game.get(gidx, []):
            raw_pname = b.get("player_name") or ""
            pname = _strip_jersey(raw_pname)
            jersey = _extract_jersey(raw_pname)
            team_idx = b.get("team_idx")
            side = b.get("team_side")
            t_var = team_id_map.get(team_idx)
            # team_idx can be wrong in some Gameone payloads; prioritize side mapping when available.
            if side == "home" and home_var != "NULL":
                t_var = home_var
            elif side == "away" and away_var != "NULL":
                t_var = away_var
            elif t_var is None and side == "home":
                t_var = home_var if home_var != "NULL" else None
            elif t_var is None and side == "away":
                t_var = away_var if away_var != "NULL" else None
            if t_var is None or not pname:
                continue
            identity_team_idx = _resolve_team_idx(team_idx, None)
            if identity_team_idx is None and side == "home":
                identity_team_idx = _resolve_team_idx(home_idx, None)
            if identity_team_idx is None and side == "away":
                identity_team_idx = _resolve_team_idx(away_idx, None)
            player_user_id = _player_user_id(pname, team_idx=identity_team_idx, team_name=None)

            # Ensure TEAM_PLAYER row exists
            out.write(
                f"INSERT INTO TEAM_PLAYER (team_id, player_id, season_id, jersey_number) "
                f"SELECT {t_var}, p.player_id, @season_id, {_num(jersey)} "
                f"FROM PLAYER p WHERE p.user_id = {_num(player_user_id)} "
                f"AND NOT EXISTS (SELECT 1 FROM TEAM_PLAYER tp2 "
                f"WHERE tp2.team_id = {t_var} AND tp2.player_id = p.player_id "
                f"AND tp2.season_id = @season_id) LIMIT 1;\n"
            )
            if jersey is not None:
                out.write(
                    f"UPDATE TEAM_PLAYER tp "
                    f"JOIN PLAYER p ON tp.player_id = p.player_id "
                    f"SET tp.jersey_number = COALESCE(tp.jersey_number, {_num(jersey)}) "
                    f"WHERE tp.team_id = {t_var} "
                    f"AND p.user_id = {_num(player_user_id)} "
                    f"AND tp.season_id = @season_id;\n"
                )
            # Ensure BATTER_STATS row exists (with correct season_type)
            _bs_extra_col = f", season_type" if stat_season_type is not None else ""
            _bs_extra_val = f", {_st_esc}" if stat_season_type is not None else ""
            out.write(
                f"INSERT INTO BATTER_STATS (tp_id, season_id{_bs_extra_col}) "
                f"SELECT tp.tp_id, @season_id{_bs_extra_val} FROM TEAM_PLAYER tp "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND tp.season_id = @season_id "
                f"AND NOT EXISTS (SELECT 1 FROM BATTER_STATS bs "
                f"WHERE bs.tp_id = tp.tp_id AND bs.season_id = @season_id "
                f"AND {_st_not_exists}) LIMIT 1;\n"
            )
            # Insert game log
            out.write(
                f"INSERT INTO BATTER_GAME_LOG "
                f"(game_idx, team_side, team_idx, player_idx, batter_stat_id, "
                f"player_name, player_position, at_bats, runs, hits, rbi, walks, strikeouts, game_type, league_code, part_code) "
                f"SELECT @game_id, {_esc(side)}, {t_var}, "
                f"(SELECT player_id FROM PLAYER WHERE user_id = {_num(player_user_id)} LIMIT 1), "
                f"(SELECT bs.batter_stat_id FROM BATTER_STATS bs "
                f"JOIN TEAM_PLAYER tp ON bs.tp_id = tp.tp_id "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND bs.season_id = @season_id AND {_st_not_exists} LIMIT 1), "
                f"{_esc(pname)}, {_esc(b.get('player_position'))}, "
                f"{_num(b.get('at_bats'))}, {_num(b.get('runs'))}, {_num(b.get('hits'))}, "
                f"{_num(b.get('rbi'))}, {_num(b.get('walks'))}, {_num(b.get('strikeouts'))}, "
                f"{_esc(game_type)}, {_esc(league_code)}, {_esc(part_code)} "
                f"FROM DUAL WHERE @game_id IS NOT NULL "
                f"AND (SELECT bs.batter_stat_id FROM BATTER_STATS bs "
                f"JOIN TEAM_PLAYER tp ON bs.tp_id = tp.tp_id "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND bs.season_id = @season_id AND {_st_not_exists} LIMIT 1) IS NOT NULL;\n"
            )

        # Pitcher game logs — ensure TEAM_PLAYER + PITCHER_STATS exist first
        for p in pitching_by_game.get(gidx, []):
            raw_pname = p.get("player_name") or ""
            pname = _strip_jersey(raw_pname)
            jersey = _extract_jersey(raw_pname)
            team_idx = p.get("team_idx")
            side = p.get("team_side")
            t_var = team_id_map.get(team_idx)
            # team_idx can be wrong in some Gameone payloads; prioritize side mapping when available.
            if side == "home" and home_var != "NULL":
                t_var = home_var
            elif side == "away" and away_var != "NULL":
                t_var = away_var
            elif t_var is None and side == "home":
                t_var = home_var if home_var != "NULL" else None
            elif t_var is None and side == "away":
                t_var = away_var if away_var != "NULL" else None
            if t_var is None or not pname:
                continue
            identity_team_idx = _resolve_team_idx(team_idx, None)
            if identity_team_idx is None and side == "home":
                identity_team_idx = _resolve_team_idx(home_idx, None)
            if identity_team_idx is None and side == "away":
                identity_team_idx = _resolve_team_idx(away_idx, None)
            player_user_id = _player_user_id(pname, team_idx=identity_team_idx, team_name=None)

            ip_val = p.get("innings_pitched")
            if ip_val is not None:
                try:
                    fval = float(ip_val)
                    whole = int(fval)
                    frac = fval - whole
                    if frac >= 0.5:
                        ip_str = f"{whole}.7"
                    elif frac >= 0.2:
                        ip_str = f"{whole}.3"
                    else:
                        ip_str = f"{whole}.0"
                except (ValueError, TypeError):
                    ip_str = "NULL"
            else:
                ip_str = "NULL"

            # Ensure TEAM_PLAYER row exists
            out.write(
                f"INSERT INTO TEAM_PLAYER (team_id, player_id, season_id, jersey_number) "
                f"SELECT {t_var}, p.player_id, @season_id, {_num(jersey)} "
                f"FROM PLAYER p WHERE p.user_id = {_num(player_user_id)} "
                f"AND NOT EXISTS (SELECT 1 FROM TEAM_PLAYER tp2 "
                f"WHERE tp2.team_id = {t_var} AND tp2.player_id = p.player_id "
                f"AND tp2.season_id = @season_id) LIMIT 1;\n"
            )
            if jersey is not None:
                out.write(
                    f"UPDATE TEAM_PLAYER tp "
                    f"JOIN PLAYER p ON tp.player_id = p.player_id "
                    f"SET tp.jersey_number = COALESCE(tp.jersey_number, {_num(jersey)}) "
                    f"WHERE tp.team_id = {t_var} "
                    f"AND p.user_id = {_num(player_user_id)} "
                    f"AND tp.season_id = @season_id;\n"
                )
            # Ensure PITCHER_STATS row exists (with correct season_type)
            _ps_extra_col = f", season_type" if stat_season_type is not None else ""
            _ps_extra_val = f", {_st_esc}" if stat_season_type is not None else ""
            out.write(
                f"INSERT INTO PITCHER_STATS (tp_id, season_id{_ps_extra_col}) "
                f"SELECT tp.tp_id, @season_id{_ps_extra_val} FROM TEAM_PLAYER tp "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND tp.season_id = @season_id "
                f"AND NOT EXISTS (SELECT 1 FROM PITCHER_STATS ps "
                f"WHERE ps.tp_id = tp.tp_id AND ps.season_id = @season_id "
                f"AND {_ps_not_exists}) LIMIT 1;\n"
            )
            # Insert game log
            out.write(
                f"INSERT INTO PITCHER_GAME_LOG "
                f"(game_idx, team_side, team_idx, player_idx, pitcher_stat_id, "
                f"player_name, innings_pitched, hits_allowed, runs_allowed, earned_runs, walks, strikeouts, game_type, league_code, part_code) "
                f"SELECT @game_id, {_esc(side)}, {t_var}, "
                f"(SELECT player_id FROM PLAYER WHERE user_id = {_num(player_user_id)} LIMIT 1), "
                f"(SELECT ps.pitcher_stat_id FROM PITCHER_STATS ps "
                f"JOIN TEAM_PLAYER tp ON ps.tp_id = tp.tp_id "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND ps.season_id = @season_id AND {_ps_not_exists} LIMIT 1), "
                f"{_esc(pname)}, {ip_str}, "
                f"{_num(p.get('hits_allowed'))}, {_num(p.get('runs_allowed'))}, "
                f"{_num(p.get('earned_runs'))}, {_num(p.get('walks'))}, {_num(p.get('strikeouts'))}, "
                f"{_esc(game_type)}, {_esc(league_code)}, {_esc(part_code)} "
                f"FROM DUAL WHERE @game_id IS NOT NULL "
                f"AND (SELECT ps.pitcher_stat_id FROM PITCHER_STATS ps "
                f"JOIN TEAM_PLAYER tp ON ps.tp_id = tp.tp_id "
                f"JOIN PLAYER p ON tp.player_id = p.player_id "
                f"WHERE p.user_id = {_num(player_user_id)} AND tp.team_id = {t_var} "
                f"AND ps.season_id = @season_id AND {_ps_not_exists} LIMIT 1) IS NOT NULL;\n"
            )

    out.write("\n")

    # 8.5. Aggregate playoff game-log data into BATTER_STATS / PITCHER_STATS
    # Playoff stats rows are created as empty containers (only tp_id, season_id, season_type).
    # These UPDATE statements fill in aggregate stats from game logs so that playoff records
    # pass isValidBatter / isValidPitcher regulation checks and become visible in the UI.
    out.write("-- 8.5 Playoff stats aggregation: fill BATTER_STATS from BATTER_GAME_LOG\n")
    out.write(
        "UPDATE BATTER_STATS bs\n"
        "JOIN (\n"
        "  SELECT bgl.batter_stat_id,\n"
        "         COUNT(DISTINCT bgl.game_idx) AS gp,\n"
        "         SUM(bgl.at_bats)             AS ab,\n"
        "         SUM(bgl.hits)                AS h,\n"
        "         SUM(bgl.rbi)                 AS rbi,\n"
        "         SUM(bgl.walks)               AS bb,\n"
        "         SUM(bgl.strikeouts)          AS so,\n"
        "         SUM(bgl.runs)                AS r\n"
        "  FROM BATTER_GAME_LOG bgl\n"
        "  GROUP BY bgl.batter_stat_id\n"
        ") agg ON bs.batter_stat_id = agg.batter_stat_id\n"
        "SET bs.games_played     = agg.gp,\n"
        "    bs.at_bats          = agg.ab,\n"
        "    bs.hits             = agg.h,\n"
        "    bs.runs_batted_in   = agg.rbi,\n"
        "    bs.walks            = agg.bb,\n"
        "    bs.strikeouts       = agg.so,\n"
        "    bs.runs_scored      = agg.r,\n"
        "    bs.plate_appearance = agg.ab + agg.bb,\n"
        "    bs.home_runs        = COALESCE(bs.home_runs, 0),\n"
        "    bs.batting_average  = CASE WHEN agg.ab > 0\n"
        "                              THEN ROUND(agg.h / agg.ab, 3)\n"
        "                              ELSE 0.000 END\n"
        "WHERE bs.season_type IS NOT NULL\n"
        "  AND bs.season_id   = @season_id\n"
        "  AND bs.batting_average IS NULL;\n"
    )
    out.write("\n")

    out.write("-- 8.5 Playoff stats aggregation: fill PITCHER_STATS from PITCHER_GAME_LOG\n")
    # innings_pitched is stored as X.0 / X.3 / X.7 (0/1/2 outs after the decimal).
    # Convert to total outs, aggregate, then convert back.
    out.write(
        "UPDATE PITCHER_STATS ps\n"
        "JOIN (\n"
        "  SELECT pgl.pitcher_stat_id,\n"
        "         COUNT(DISTINCT pgl.game_idx) AS gp,\n"
        "         COALESCE(SUM(\n"
        "           FLOOR(IFNULL(pgl.innings_pitched, 0)) * 3 +\n"
        "           CASE\n"
        "             WHEN IFNULL(pgl.innings_pitched, 0)\n"
        "                  - FLOOR(IFNULL(pgl.innings_pitched, 0)) >= 0.6 THEN 2\n"
        "             WHEN IFNULL(pgl.innings_pitched, 0)\n"
        "                  - FLOOR(IFNULL(pgl.innings_pitched, 0)) >= 0.2 THEN 1\n"
        "             ELSE 0\n"
        "           END\n"
        "         ), 0)                        AS total_outs,\n"
        "         SUM(pgl.strikeouts)          AS so,\n"
        "         SUM(pgl.walks)               AS bb,\n"
        "         SUM(pgl.hits_allowed)        AS h,\n"
        "         SUM(pgl.earned_runs)         AS er\n"
        "  FROM PITCHER_GAME_LOG pgl\n"
        "  GROUP BY pgl.pitcher_stat_id\n"
        ") agg ON ps.pitcher_stat_id = agg.pitcher_stat_id\n"
        "SET ps.games_played    = agg.gp,\n"
        "    ps.innings_pitched = FLOOR(agg.total_outs / 3) +\n"
        "                         CASE agg.total_outs % 3\n"
        "                           WHEN 0 THEN 0.0\n"
        "                           WHEN 1 THEN 0.3\n"
        "                           ELSE       0.7\n"
        "                         END,\n"
        "    ps.strikeouts      = agg.so,\n"
        "    ps.walks_allowed   = agg.bb,\n"
        "    ps.wins            = COALESCE(ps.wins,   0),\n"
        "    ps.losses          = COALESCE(ps.losses, 0),\n"
        "    ps.saves           = COALESCE(ps.saves,  0),\n"
        "    ps.era             = CASE WHEN agg.total_outs > 0\n"
        "                              THEN ROUND((agg.er * 27.0) / agg.total_outs, 2)\n"
        "                              ELSE 0.00 END,\n"
        "    ps.whip            = CASE WHEN agg.total_outs > 0\n"
        "                              THEN ROUND(((agg.h + agg.bb) * 3.0) / agg.total_outs, 2)\n"
        "                              ELSE 0.00 END\n"
        "WHERE ps.season_type IS NOT NULL\n"
        "  AND ps.season_id   = @season_id\n"
        "  AND ps.era IS NULL;\n"
    )
    out.write("\n")

    # 8. TEAM_PLAYER.part_code — backfill from accumulated league record context
    # Each team's most-common part_code from league records is propagated to TEAM_PLAYER.
    if team_context_counter:
        out.write("-- Team-Player part_code backfill from league records\n")
        for norm_name, bucket in sorted(team_context_counter.items()):
            part_counter = bucket["part"]
            if not part_counter:
                continue
            best_part_code = part_counter.most_common(1)[0][0]
            team_var = team_name_norm_map.get(norm_name)
            if team_var is None:
                continue
            out.write(
                f"UPDATE TEAM_PLAYER SET part_code = {_esc(best_part_code)} "
                f"WHERE team_id = {team_var} "
                f"AND season_id = @season_id "
                f"AND part_code IS NULL;\n"
            )
        out.write("\n")

    out.write("SET FOREIGN_KEY_CHECKS = 1;\n")
    out.write("-- Done.\n")


if __name__ == "__main__":
    main()
