"""Build a Korean audit report from the administrator's allowlisted JSON export.

Run with the bundled document runtime. This script makes no network requests and
never edits the input, a sync run, player records, or a published revision.
"""
import argparse
import collections
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = 'Pretendard Variable'

SOURCE_RBI_CHECKS = {
    '55840': ('서경대학교 JUKSITA', 7, 10, '원본에서도 타점 10과 득점 7이 표시됨'),
    '55783': ('한신대학교 갱스터', 1, 0, '원본도 타점 0이며 득점 경위 판정 필요'),
    '55558': ('서강대학교 알바트로스', 10, 0, '한정우 3회 좌월홈에도 개인 타점과 득점 0'),
    '55214': ('서울시립대학교 FALCONS', 5, 0, '박휴준 1회 중월홈에도 개인 타점과 득점 0'),
    '55133': ('외대(글로벌) 유니온', 7, 0, '원본도 타점 0이며 득점 경위 판정 필요'),
    '54987': ('경희대학교(서울) BRAVES', 10, 0, '원본에서도 선수 득점과 타점 전원 0'),
    '54897': ('인천대학교 바이킹', 2, 0, '원본도 타점 0이며 득점 경위 판정 필요'),
}
ISSUE_LABELS = {
    'DETAIL_PITCHER_HITS_TOTAL': '투수 피안타 합계 불일치',
    'DETAIL_PITCHER_RUNS_TOTAL': '투수 실점 합계 불일치',
    'DETAIL_PITCHER_EARNED_RUNS': '투수 자책점이 실점 초과',
    'DETAIL_BATTER_TOTAL': '타자 득점 합계 불일치',
    'DETAIL_TEAM_RBI': '타점 합계가 팀 득점 초과',
    'DETAIL_TEAM_RBI_ZERO': '득점 팀의 타점 합계 0',
    'DETAIL_LINE_SCORE': '이닝 점수 확인 필요',
}


def finite_sum(rows, field):
    values = [r.get('stats', {}).get(field) for r in rows]
    return sum(values) if values and all(isinstance(v, (int, float)) for v in values) else None


def summarize(data):
    games = []
    for item in sorted(data['games'], key=lambda g: (g['game']['playedAt'], g['sourceGameId'])):
        detail = item.get('detail') or {}
        game = dict(item['game'])
        game.update(sourceGameId=item['sourceGameId'], providerGameId=detail.get('providerGameId'),
                    detailStatus=detail.get('status'), quality=item.get('quality'), issues=[], teams=[])
        assert re.fullmatch(r'\d{1,24}', game['providerGameId'] or ''), 'Missing provider URL identifier'
        game['sourceUrl'] = f"https://unique-play.com/game/{game['providerGameId']}/boxscore"
        for team in detail.get('teams', []):
            game['teams'].append({
                'teamName': team['teamName'], 'totals': team['totals'],
                'batterRows': len(team['batters']), 'pitcherRows': len(team['pitchers']),
                'atBats': finite_sum(team['batters'], 'atBats'),
                'hits': finite_sum(team['batters'], 'hits'),
                'runs': finite_sum(team['batters'], 'runs'), 'rbi': finite_sum(team['batters'], 'rbi'),
                'pitcherHits': finite_sum(team['pitchers'], 'hitsAllowed'),
                'pitcherRuns': finite_sum(team['pitchers'], 'runsAllowed'),
            })
        for issue in item.get('issues', []):
            clean = dict(issue)
            clean['playerName'] = next((r['playerName'] for t in detail.get('teams', [])
                                       for r in t['batters'] + t['pitchers'] if r['rowKey'] == issue.get('rowKey')), None)
            game['issues'].append(clean)
        game['homeRunChecks'] = []
        for team in detail.get('teams', []):
            for batter in team['batters']:
                hrs = [p for p in batter['plateAppearances'] if re.search(r'홈', p.get('result') or '')]
                if hrs and any(batter['stats'].get(k) is not None and batter['stats'][k] < len(hrs) for k in ('rbi', 'runs', 'hits')):
                    game['homeRunChecks'].append({'teamName': team['teamName'], 'playerName': batter['playerName'],
                                                 'plateAppearances': hrs, 'rbi': batter['stats'].get('rbi'),
                                                 'runs': batter['stats'].get('runs')})
        game['rbiSourceVerified'] = game['providerGameId'] in SOURCE_RBI_CHECKS
        games.append(game)
    assert len(games) == 117 and len({g['sourceGameId'] for g in games}) == 117
    assert len({g['providerGameId'] for g in games}) == 117
    counts = collections.Counter(i['code'] for g in games for i in g['issues'])
    return {
        'schemaVersion': 1, 'auditDateKst': '2026-09-05',
        'scope': '저장된 수집 후보 전수 산술검사와 타점 7경기 원본 대조',
        'runId': data['runId'], 'checksum': data['checksum'], 'reviewChecksum': data['reviewChecksum'],
        'expectedRevision': data['expectedRevision'], 'issueCounts': dict(counts),
        'totalGames': len(games), 'availableGames': sum(g['detailStatus'] == 'AVAILABLE' for g in games),
        'warningGames': sum(bool(g['issues']) for g in games),
        'warningOccurrences': sum(counts.values()), 'games': games,
    }


def set_font(run, size=11, bold=False, color='000000'):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        run._element.get_or_add_rPr().rFonts.set(qn('w:' + attr), FONT)


def paragraph(doc, text, style=None):
    previous = doc.element.body[-2] if len(doc.element.body) > 1 else None
    p = doc.add_paragraph(style=style)
    if previous is not None and previous.tag == qn('w:tbl'):
        p.paragraph_format.space_before = Pt(8)
    p.add_run(text)
    return p


def link(p, label, url, size=10):
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    for tag, val in [('color', '183C79'), ('u', 'single'), ('sz', str(int(size * 2)))]:
        el = OxmlElement('w:' + tag); el.set(qn('w:val'), val); props.append(el)
    r.append(props); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def table(doc, headers, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: el.set(qn('w:' + key), val)
        borders.append(el)
    tblpr.append(borders)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    fill_row(t.rows[0], headers, widths, header=True)
    return t


def fill_row(row, texts, widths, header=False, shaded=False):
    no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    for idx, (cell, text, width) in enumerate(zip(row.cells, texts, widths)):
        cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
        for side, size in [('top', '90'), ('bottom', '90'), ('left', '85'), ('right', '85')]:
            el = OxmlElement('w:' + side); el.set(qn('w:w'), size); el.set(qn('w:type'), 'dxa'); margins.append(el)
        pr.append(margins)
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '162A53' if header else 'F2F5FA' if shaded else 'FFFFFF'); pr.append(shade)
        cell.text = str(text)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs: set_font(run, 10, header, 'FFFFFF' if header else '000000')


def issue_text(i):
    a, b = i.get('observed'), i.get('expected')
    code = i['code']
    label = {
        'DETAIL_PITCHER_HITS_TOTAL': f'투수 피안타 {a} / 상대 안타 {b}',
        'DETAIL_PITCHER_RUNS_TOTAL': f'투수 실점 {a} / 상대 득점 {b}',
        'DETAIL_PITCHER_EARNED_RUNS': f"{i.get('playerName') or '투수'} 자책 {a} / 실점 {b}",
        'DETAIL_BATTER_TOTAL': f'타자 득점 합계 {a} / 팀 득점 {b}',
        'DETAIL_TEAM_RBI': f'타점 합계 {a} / 팀 득점 {b}',
        'DETAIL_TEAM_RBI_ZERO': f'타점 합계 {a} / 팀 득점 {b} — 경위 확인',
        'DETAIL_LINE_SCORE': f'이닝 점수 미기입 또는 합계 확인 / 팀 득점 {b}',
    }.get(code, i['message'])
    return label


def build_doc(summary, output):
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.header_distance = sec.footer_distance = Inches(.28)
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    normal = doc.styles['Normal']; normal.font.name = FONT; normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.17
    for name, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 16), ('Heading 2', 12)]:
        s = doc.styles[name]; s.font.name = FONT; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0, 0, 0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        s.paragraph_format.space_before = Pt(14 if name.startswith('Heading') else 0)
        s.paragraph_format.space_after = Pt(7); s.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]; set_font(header.add_run('AUBL 2026 기록 검수 보고서'), 9)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run('2026년 9월 5일  |  '), 9)
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)
    doc.core_properties.title = '2026 AUBL 경기 기록 전수 검수 보고서'
    doc.core_properties.subject = '저장된 117경기 후보와 타점 7경기 원본 대조'
    doc.core_properties.author = 'AUBL 기록 검수 작업'; doc.core_properties.last_modified_by = 'AUBL 기록 검수 작업'
    doc.core_properties.comments = ''

    paragraph(doc, '2026 AUBL 경기 기록\n전수 검수 보고서', 'Title')
    paragraph(doc, '검수일 2026년 9월 5일    대상 AUBL 운영 관리자', 'Subtitle')
    paragraph(doc, '저장된 UniquePlay 수집 후보 117경기를 전수 검사한 결과, 상세 기록이 있는 112경기 중 46경기에서 80건의 검토 경고가 발견됐다. 66경기는 이번 검사에서 경고가 없었고, 5경기는 원천 상세 미게시로 선수 기록을 검사할 수 없었다. 전체 117경기의 날짜, 경기명, 원본 링크, 대상 팀과 불일치 수치는 이 문서의 전체 목록에 수록했다.')
    paragraph(doc, '타점 관련 7경기는 UniquePlay 원본 화면과 추가로 대조했다. 원본도 수집값과 같아 이 7건을 단순히 AUBL 표시 과정의 누락으로 볼 수는 없다. 다만 득점이 있어도 실책 등으로 타점이 없을 수 있으므로 타점 0 자체를 오류로 확정하거나 득점과 같게 보정하지 않는다.')
    paragraph(doc, '1 검사 범위와 결과', 'Heading 1')
    widths = [.95, 1.0, 5.15]
    t = table(doc, ['구분', '경기 수', '판단 범위'], widths)
    for n, row in enumerate([
        ['전체', '117', '저장된 수집 후보의 경기 ID 117개 및 원본 링크 117개 중복 없음'],
        ['검토 필요', '46', '80개 경고 발생. 여러 경고가 있는 경기는 1경기로 집계'],
        ['경고 없음', '66', '검사한 필드의 산술 조건 통과. 실제 경기 전체의 정상 판정은 아님'],
        ['검사 불가', '5', '경기 결과는 있지만 상세 미게시로 타자와 투수 기록 미확인'],
    ]): fill_row(t.add_row(), row, widths, shaded=n % 2 == 1)
    paragraph(doc, '검사 대상은 2026년 3월 21일부터 8월 24일까지의 종료 경기 후보다. 검사 기준은 팀 득점과 이닝·타자 득점 합계, 타점 범위, 타수 대비 안타, 투수 실점·피안타와 상대 기록, 투수 자책점과 실점이다. 타점 누락은 null과 0을 구분해 검사했다. 홈런 결과와 선수 타점·득점의 모순도 별도로 점검했다.')
    paragraph(doc, '조사 한계', 'Heading 2')
    paragraph(doc, '전수 검사는 2026년 9월 5일 13시 46분경 완료된 v9 수집 후보를 기준으로 한다. 117경기 원본 페이지를 같은 시각에 모두 다시 열어 검증한 것은 아니다. 이번 원본 추가 대조 범위는 타점 7경기이며, 모든 투구와 주자 이동 로그는 수집되지 않아 득점 경위를 완전히 재구성할 수 없다. 경기별 원본 확인과 공식 기록원의 판정이 더 필요한 항목을 구분해 남겼다.')

    paragraph(doc, '2 경고 유형별 집계', 'Heading 1')
    widths = [.5, 2.85, 3.75]; t = table(doc, ['건수', '검사 항목', '해석과 우선 조치'], widths)
    notes = {
        'DETAIL_PITCHER_HITS_TOTAL': '양 팀 안타와 모든 투수 피안타를 대조한다.',
        'DETAIL_PITCHER_RUNS_TOTAL': '상대 팀 득점과 투수별 실점 귀속을 확인한다.',
        'DETAIL_PITCHER_EARNED_RUNS': '18개 선수 행에서 검출. 자책점과 실점의 열 및 원천 값을 확인한다.',
        'DETAIL_BATTER_TOTAL': '팀 득점과 개인 득점의 미기입 또는 합계 차이를 확인한다.',
        'DETAIL_TEAM_RBI': '서경대 7득점에 타점 10. 원본에서도 같은 수치 확인.',
        'DETAIL_TEAM_RBI_ZERO': '실책 등으로 가능하므로 자동 오류 확정 금지. 홈런 동반 2경기는 우선 확인.',
        'DETAIL_LINE_SCORE': '강남대 이닝 점수 미기입. 최종 0점만으로 전 이닝 0을 채우지 않는다.',
    }
    for n, (code, count) in enumerate(summary['issueCounts'].items()):
        fill_row(t.add_row(), [count, ISSUE_LABELS.get(code, code), notes[code]], widths, shaded=n % 2 == 1)
    paragraph(doc, '80건은 검사 경고 수이며 선수 수나 경기 수가 아니다. 타점 7건은 이 집계 안에 포함된다. 홈런과 모순되는 2건도 기존 46경기에 포함되므로 경기 수에 다시 더하지 않았다.')

    paragraph(doc, '3 타점 원본 대조 결과', 'Heading 1')
    paragraph(doc, '다음 표의 날짜는 KST다. 원본 타자 표의 개인 타점·득점과 합계가 저장된 후보와 같은 것을 확인했다. 원본의 같은 값은 전송·표시 누락 여부를 판단하는 근거이며, 해당 값이 실제 경기에서 올바르다는 보증은 아니다.')
    widths = [.8, 2.2, .8, 3.3]; t = table(doc, ['날짜', '대상 팀', '득점\n타점', '원본 확인과 권고'], widths)
    for n, g in enumerate(g for g in summary['games'] if g['rbiSourceVerified']):
        team, runs, rbi, note = SOURCE_RBI_CHECKS[g['providerGameId']]
        row = t.add_row(); fill_row(row, [g['playedAt'][5:10], team, f'{runs} / {rbi}', note], widths, shaded=n % 2 == 1)
        p = row.cells[3].add_paragraph(); link(p, '원본 ' + g['providerGameId'], g['sourceUrl'])
    paragraph(doc, '서강대 한정우는 3회 좌월홈, 시립대 박휴준은 1회 중월홈으로 기록돼 있는데 각각 개인 타점과 득점이 0이다. 홈런 결과와 누적된 개인 수치가 서로 맞지 않아 공식 기록 확인이 우선 필요하다. 경희대도 팀 10득점에 선수 득점이 전부 0이므로 타점과 함께 검수한다.')

    paragraph(doc, '4 선수 삭제 후보가 표시되는 이유', 'Heading 1')
    paragraph(doc, '현재 실행의 삭제 후보 419건은 실제 선수 419명의 삭제가 아니다. 비교기는 기존 공개 누적 기록의 키가 새 후보에서 발견되지 않으면 DELETE로 분류한다. 같은 선수가 타자·투수 기록을 모두 갖는 경우도 있으므로 이 숫자를 고유 선수 수로 읽으면 안 된다.')
    paragraph(doc, '팀 별칭과 비교 키', 'Heading 2')
    paragraph(doc, '기존 데이터는 AUBL DB 팀명, 후보는 UniquePlay 원천 팀명을 사용한다. 저장된 팀 연결을 이 비교 키에 적용하지 않아 같은 선수도 신규와 삭제가 한 쌍으로 나타날 수 있다. 예를 들어 같은 팀으로 연결된 “홈대 야구부”와 “홈대학교”라도 비교 문자열은 다르다. 이 예시는 설명용이며 실제 선수의 삭제 사례를 뜻하지 않는다.')
    paragraph(doc, '규정 IN OUT 이동', 'Heading 2')
    paragraph(doc, '비교 키에 규정 타석·투구이닝 충족 여부인 IN/OUT이 들어 있다. 선수가 OUT에서 IN으로 이동하면 현재 비교기는 변경 1건 대신 삭제와 신규로 분리할 수 있다. 반면 게시기는 연결된 팀 선수·시즌·제공자로 같은 누적 행을 갱신하므로 비교 방식과 실제 갱신 기준이 일치하지 않는다.')
    paragraph(doc, '수집 완전성과 실제 게시 영향', 'Heading 2')
    paragraph(doc, '기존 가상 스크롤 수집에는 마지막 행에 도달하기 전에 종료될 가능성이 있었다. v10은 끝까지 이동한 뒤 3회 연속 행이 안정될 때 완료로 판단하고 최대 반복수 초과는 실패로 처리한다. 그러나 새 워커를 배포했다고 이전 v9 후보가 자동으로 보완되지는 않는다.')
    paragraph(doc, '게시기는 Player 또는 TeamPlayer 자체를 삭제하지 않는다. 실제로 새 후보에서 빠진 누적 기록은 이전 리비전에 남고 새 공개 리비전의 집계에서 제외될 수 있다. 별칭으로만 다르게 인식된 행은 게시 과정에서 같은 행을 갱신할 수도 있다. 따라서 419건 전체를 삭제 승인하지 말고, 비교 키를 통일한 재대조로 신규·변경·실제 누락을 나눠야 한다. 이 보고서에서 419건의 원인별 개수는 확정하지 않았다.')

    paragraph(doc, '5 관리자 처리와 배포 상태', 'Heading 1')
    paragraph(doc, '치명적이지 않은 상세 합계 경고는 관리자가 경고 내용을 확인한 뒤 게시할 수 있도록 변경했다. 다만 팀·경기 식별 오류, 점수 구조 불일치, 개인정보 또는 형식 위반 등 차단 오류는 계속 게시를 막는다. 공개 상세 및 선수 기록에는 오류 수정 중 상태와 이후 해결 상태를 표시하는 코드를 추가했다.')
    paragraph(doc, '관리자는 동기화 검수의 경기 상세에서 원본과 선수 기록을 확인하고, 수정할 필드와 사유를 입력한다. 수정은 수집 원본과 분리해 보존하며 게시 전까지 공개 데이터는 바뀌지 않는다. 이미 게시된 자료는 수정 후보를 만들어 같은 검수 절차를 밟는다. 다음 수동 동기화에서 원천 값이 수정값과 일치하면 원천 해결을 알리고, 서로 다르면 관리자 선택이 필요한 충돌로 남긴다.')
    paragraph(doc, 'Docker Hub 업로드 승인 후 백엔드 v25와 워커 v10을 NAS에 배포했다. 백엔드는 15시 00분, 워커는 15시 18분에 정상 기동을 확인했다. MariaDB 스키마 버전 6은 그대로이며 추가 마이그레이션은 실행되지 않았다. 웹 변경은 로컬 소스, Flutter 변경은 앱 소스에 반영했으며 이 작업에서 웹 호스팅이나 앱 스토어 배포는 하지 않았다.')
    paragraph(doc, '이번 조사에서는 새 수집, 실제 선수 수치 수정, 공개 리비전 게시·활성화를 실행하지 않았다. 기존 공개 리비전을 유지했다. 특히 선수 삭제 후보의 비교 문제는 남아 있으므로 상세 경고를 게시 가능하게 바꾼 것과 별개로 현재 후보의 게시를 보류한다.')
    paragraph(doc, '검증과 남은 확인', 'Heading 2')
    paragraph(doc, '백엔드 106개, 워커와 웹 계약·검수 테스트 120개, Flutter 143개 테스트가 통과했다. 관리자 다운로드의 비공개 메모 제외와 표의 null·0 표시도 회귀 테스트에 포함했다. 실제 오류 수정 후 게시·복구 절차는 운영 데이터에 쓰기 없이 로컬 테스트로 확인했으며, 운영에서의 최초 게시 검증은 별도로 남겨 둔다.')

    doc.add_page_break()
    paragraph(doc, '6 전체 117경기 검수 목록', 'Heading 1')
    paragraph(doc, '경기 날짜 오름차순이다. 경기 열의 위 팀과 아래 팀 순서로 스코어를 표시한다. 대상 팀별 수치는 수집 후보의 값이다. “경고 없음”은 이번 검사에서 모순이 검출되지 않았다는 뜻이며, “상세 미게시”는 선수 기록 검사 대상에서 제외됐다. 링크는 UniquePlay 원본 박스스코어로 연결된다. 상세 미게시 경기의 7대0 결과도 자료만으로 몰수승으로 단정하지 않는다.')
    widths = [.35, 1.1, 2.05, 3.6]
    t = table(doc, ['번호', '경기 일시\n조와 구장', '경기명과 결과\n원본 링크', '대상 팀과 불일치 내용'], widths)
    for n, g in enumerate(summary['games'], 1):
        d = g['playedAt']; date = f"{d[5:10]} {d[11:16]}\n{g['groupCode']}조\n{g['venue']}"
        matchup = f"{g['homeTeamName']}\n{g['homeScore']} : {g['awayScore']}\n{g['awayTeamName']}"
        if g['detailStatus'] != 'AVAILABLE':
            result = '상세 미게시\n선수 기록 검사 불가. 결과만 확인됨.'
        elif not g['issues']:
            result = '경고 없음\n이번 산술검사 기준 통과.'
        else:
            grouped = collections.defaultdict(list)
            for issue in g['issues']: grouped[issue.get('teamName') or '경기'].append(issue_text(issue))
            result = '검토 필요\n' + '\n'.join(team + '\n' + '\n'.join(lines) for team, lines in grouped.items())
            for h in g['homeRunChecks']:
                result += f"\n추가 확인 {h['playerName']} {h['plateAppearances'][0]['inning']}회 홈런 / 타점 {h['rbi']} / 득점 {h['runs']}"
            if g['rbiSourceVerified']: result += '\n타점 수치는 원본 대조 완료'
        row = t.add_row(); fill_row(row, [n, date, matchup, result], widths, shaded=n % 2 == 0)
        p = row.cells[2].add_paragraph(); p.paragraph_format.space_after = Pt(0)
        link(p, 'UniquePlay ' + g['providerGameId'], g['sourceUrl'])
    doc.add_page_break()
    paragraph(doc, '7 검수 자료 식별 정보', 'Heading 1')
    paragraph(doc, '아래 식별값은 어떤 수집 후보를 조사했는지 재현하기 위한 값이며 로그인 토큰이나 비밀번호가 아니다. 관리자 메모, 계정 식별정보, 세션 정보와 원문 HTML은 보고서에 포함하지 않았다.')
    for label, key in [('수집 실행', 'runId'), ('수집 후보 체크섬', 'checksum'), ('검수 체크섬', 'reviewChecksum'), ('기준 공개 리비전', 'expectedRevision')]:
        p = paragraph(doc, label + '\n' + summary[key])
        for r in p.runs: set_font(r, 9.5)
    paragraph(doc, '근거는 관리자 경기 검수 API의 허용 필드 내보내기, 해당 후보의 117경기 상세 표, 타점 7경기의 UniquePlay 원본 표, 백엔드 비교·게시 코드 및 NAS 기동 결과다. 게시 여부와 원천 수치는 조사 이후 달라질 수 있으므로 재검수 시 실행 ID와 체크섬을 함께 기록한다.')
    doc.save(output)


def main():
    p = argparse.ArgumentParser(); p.add_argument('input'); p.add_argument('--output', required=True); p.add_argument('--summary', required=True)
    args = p.parse_args(); source = Path(args.input)
    data = json.loads(source.read_text()); result = summarize(data)
    result['inputSha256'] = hashlib.sha256(source.read_bytes()).hexdigest()
    for path in (Path(args.output), Path(args.summary)): path.parent.mkdir(parents=True, exist_ok=True)
    Path(args.summary).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    build_doc(result, args.output)
    print(json.dumps({k: result[k] for k in ['totalGames', 'availableGames', 'warningGames', 'warningOccurrences', 'issueCounts']}, ensure_ascii=False))
    print(args.output)


if __name__ == '__main__':
    main()
